"""
build_trajectory_model.py
──────────────────────────
Trajectory Model — age-curve development & projection scouting model.

Every existing model in this repo (Lamberts Index, WAR, Hockey-Style, Wyscout
composite scores) rates a player on CURRENT output. This model instead asks:
where is this player in their development curve, and what should we expect
from them in 1 / 3 / 5 years?

Methodology
-----------
1. Current Level  — position-specific weighted percentile score (0-100),
   same style as the Lamberts SQS, built from each position group's own
   metric blueprint.

2. Age Curve       — per position group, fit a minutes-weighted quadratic
   (Current Level ~ a·Age² + b·Age + c) across the full multi-league pool.
   The vertex of the parabola gives the population's expected Peak Age.
   This is a CROSS-SECTIONAL curve (one snapshot per player, no season-over-
   season panel data is available in the Wyscout exports) — it describes how
   players of different ages currently compare, which is the standard
   fallback when longitudinal history isn't available. Treat it as a
   population baseline, not a guarantee for any individual.

3. Vs Age Curve    — a player's residual = Current Level − curve(Age).
   Positive = performing above what's typical for their age; negative =
   below. This residual is assumed to persist and is carried forward onto
   the curve at future ages to build the projection (the standard
   "delta method" simplification).

4. Projections     — Proj(Age+k) = clip(curve(Age+k) + residual, 1, 99)
   for k = 1, 3, 5 years.

5. Trajectory Score = 0.35 × Current Level + 0.65 × Proj(+3y)
   Weighted toward the 3-year horizon since recruitment decisions are made
   for the players a club will have on the books over a typical contract,
   not just the season they're bought in.

6. Trend label, from age relative to the position's Peak Age:
   ASCENDING → APPROACHING PEAK → PEAK WINDOW → EARLY DECLINE → LATE DECLINE

Output: reports/Trajectory_Model.xlsx
  README · Age Curves · All Players · GK/CB/FB/DM/CM/W/FW ·
  Breakout Prospects · Decline Risk

Usage:
  python build_trajectory_model.py
  python build_trajectory_model.py --leagues "Czech II" Slovakia --min-minutes 700
  python build_trajectory_model.py --output reports/My_Trajectory_Model.xlsx
"""
from __future__ import annotations

import argparse
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

# ── Config ─────────────────────────────────────────────────────────────────────

ROOT = Path(__file__).parent
WYSCOUT_DIR = ROOT / "Wyscout Files"
OUT_DIR = ROOT / "reports"

SKIP_FILES = {"FCHK Model V3 - Loaded Leagues", "FCHK Model V3 - Model Input",
              "FCHK Model V3 - Player Scores", "FCHK Model V3 - Player Styles",
              "FCHK Model V3 - Recruitment Scores", "FCHK Model V3 - Smart Club Closeness",
              "FCHK Model V3 - Summary", "FCHK Model V3 Scores", "FCHK Scouting Report",
              "Leagues Overview", "Wyscout Anomaly Report", "Wyscout Full Scouting Report"}

DEFAULT_LEAGUES = None  # None = all leagues in WYSCOUT_DIR
DEFAULT_MIN_MINUTES = 500
DEFAULT_MIN_FIT_N = 50   # min players required to trust a fitted curve over the fallback

# Position mapping (first token of Wyscout position string → group)
POS_MAP: dict[str, str] = {
    "CF": "FW", "SS": "FW",
    "LW": "W",  "RW": "W",  "LWF": "W",  "RWF": "W",  "WF": "W",
    "LAMF": "W", "RAMF": "W",
    "AMF": "CM",
    "CMF": "CM", "LCM": "CM", "RCM": "CM", "LCMF": "CM", "RCMF": "CM",
    "DMF": "DM", "LDM": "DM", "RDM": "DM", "LDMF": "DM", "RDMF": "DM",
    "LB": "FB",  "RB": "FB",  "LWB": "FB", "RWB": "FB",
    "CB": "CB",  "LCB": "CB", "RCB": "CB",
    "GK": "GK",
}

# Population fallback peak ages, used only when a position group has too few
# players to fit a reliable curve, or when the fit doesn't produce a concave
# (rise-then-fall) shape.
DEFAULT_PEAK_AGE: dict[str, float] = {
    "GK": 29.0, "CB": 28.0, "FB": 27.0, "DM": 27.0, "CM": 26.0, "W": 25.0, "FW": 25.0,
}

# Per-position Current Level metrics (weights) — mirrors the SQS blueprint style
LEVEL_BLUEPRINTS: dict[str, list[tuple[str, float]]] = {
    "GK": [
        ("Save rate, %", 4.0),
        ("Prevented goals per 90", 3.0),
        ("Exits per 90", 2.0),
        ("Aerial duels per 90", 1.5),
        ("Accurate passes, %", 1.5),
        ("Accurate long passes, %", 1.0),
    ],
    "CB": [
        ("Successful defensive actions per 90", 3.0),
        ("Defensive duels won, %", 2.5),
        ("Aerial duels won, %", 2.0),
        ("Interceptions per 90", 2.0),
        ("PAdj Interceptions", 1.5),
        ("Shots blocked per 90", 1.0),
        ("Accurate passes, %", 1.0),
        ("Progressive passes per 90", 1.0),
    ],
    "FB": [
        ("Crosses per 90", 2.0),
        ("Accurate crosses, %", 1.5),
        ("xA per 90", 2.0),
        ("Assists per 90", 1.5),
        ("Progressive runs per 90", 1.5),
        ("Dribbles per 90", 1.0),
        ("Successful defensive actions per 90", 2.0),
        ("Defensive duels won, %", 2.0),
        ("Aerial duels won, %", 1.0),
        ("Progressive passes per 90", 1.5),
    ],
    "DM": [
        ("Successful defensive actions per 90", 3.0),
        ("Defensive duels won, %", 2.5),
        ("Interceptions per 90", 2.5),
        ("PAdj Interceptions", 2.0),
        ("Aerial duels won, %", 1.5),
        ("Passes per 90", 1.5),
        ("Accurate passes, %", 1.5),
        ("Progressive passes per 90", 1.5),
    ],
    "CM": [
        ("Passes per 90", 2.0),
        ("Accurate passes, %", 2.0),
        ("Progressive passes per 90", 2.5),
        ("Key passes per 90", 2.5),
        ("xA per 90", 2.0),
        ("Assists per 90", 1.5),
        ("Progressive runs per 90", 1.5),
        ("Successful defensive actions per 90", 1.5),
        ("Goals per 90", 1.5),
    ],
    "W": [
        ("Goals per 90", 2.5),
        ("xG per 90", 2.0),
        ("Assists per 90", 2.0),
        ("xA per 90", 2.0),
        ("Dribbles per 90", 2.0),
        ("Successful dribbles, %", 1.5),
        ("Progressive runs per 90", 2.0),
        ("Touches in box per 90", 2.0),
        ("Key passes per 90", 1.5),
    ],
    "FW": [
        ("Goals per 90", 4.0),
        ("xG per 90", 3.0),
        ("Non-penalty goals per 90", 2.5),
        ("Shots per 90", 1.5),
        ("Shots on target, %", 1.5),
        ("Goal conversion, %", 1.5),
        ("Touches in box per 90", 1.5),
        ("Aerial duels won, %", 1.5),
        ("Dribbles per 90", 1.0),
        ("xA per 90", 1.0),
    ],
}

# ── Colors ─────────────────────────────────────────────────────────────────────
C = {
    "navy":    "0D1B2A",
    "gold":    "C9A84C",
    "header":  "154360",
    "light":   "EBF5FB",
    "white":   "FFFFFF",
    "breakout":   "1A5276",
    "highceil":   "117A65",
    "ascending":  "1E8449",
    "stable":     "626567",
    "decline":    "922B21",
    "trend_asc":  "1E8449",
    "trend_appr": "117A65",
    "trend_peak": "B7950B",
    "trend_early":"935116",
    "trend_late": "922B21",
}


# ── Data loading ───────────────────────────────────────────────────────────────

def load_leagues(leagues: list[str] | None, min_minutes: int) -> pd.DataFrame:
    if leagues is None:
        paths = sorted(p for p in WYSCOUT_DIR.glob("*.xlsx") if p.stem not in SKIP_FILES)
    else:
        paths = [WYSCOUT_DIR / f"{lg}.xlsx" for lg in leagues]

    frames: list[pd.DataFrame] = []
    for path in paths:
        if not path.exists():
            print(f"  [warn] {path} not found — skipping")
            continue
        lg = path.stem
        try:
            df = pd.read_excel(path)
        except Exception as e:
            print(f"  [warn] Could not read {path.name}: {e}")
            continue
        df = df.copy()
        df["_League"] = lg
        frames.append(df)

    if not frames:
        raise RuntimeError(f"No Wyscout files found in {WYSCOUT_DIR}")

    raw = pd.concat(frames, ignore_index=True)
    print(f"  Loaded {len(frames)} league files, {len(raw)} player rows")

    mins_col = next((c for c in ["Minutes played", "MinutesPlayed", "Minutes"] if c in raw.columns), None)
    raw["_minutes"] = pd.to_numeric(raw[mins_col], errors="coerce").fillna(0) if mins_col else 0

    raw = raw[raw["_minutes"] >= min_minutes].copy()
    print(f"  → {len(raw)} players after {min_minutes}+ minute filter")
    return raw.reset_index(drop=True)


def map_position(pos_str: str) -> str:
    if not isinstance(pos_str, str):
        return "Other"
    first = pos_str.split(",")[0].strip()
    return POS_MAP.get(first, "Other")


def add_position_group(df: pd.DataFrame) -> pd.DataFrame:
    pos_col = next((c for c in ["Position", "Pos"] if c in df.columns), None)
    if pos_col:
        df["_pos_group"] = df[pos_col].apply(map_position)
        df["_full_position"] = df[pos_col].fillna("Unknown")
    else:
        df["_pos_group"] = "Other"
        df["_full_position"] = "Unknown"
    return df


# ── Current Level score ─────────────────────────────────────────────────────────

def compute_current_level(df: pd.DataFrame) -> pd.DataFrame:
    """Position-specific weighted percentile score (0-100) — 'where they stand today'."""
    df = df.copy()
    df["_level"] = np.nan

    for pos, blueprint in LEVEL_BLUEPRINTS.items():
        mask = df["_pos_group"] == pos
        if mask.sum() == 0:
            continue
        grp = df.loc[mask]
        score = pd.Series(0.0, index=grp.index)
        total_w = 0.0
        for metric, w in blueprint:
            if metric in grp.columns:
                vals = pd.to_numeric(grp[metric], errors="coerce").fillna(0)
                pct = vals.rank(pct=True) * 100
                score += w * pct
                total_w += w
        if total_w > 0:
            score /= total_w
        df.loc[mask, "_level"] = score.values

    return df


# ── Age curve fitting ────────────────────────────────────────────────────────────

def fit_age_curve(sub: pd.DataFrame, pos: str) -> dict:
    """Minutes-weighted quadratic fit of Current Level vs Age for one position group."""
    ages = pd.to_numeric(sub["Age"], errors="coerce")
    levels = sub["_level"]
    mins = sub["_minutes"]

    mask = ages.notna() & levels.notna() & (ages >= 15) & (ages <= 42)
    ages_v = ages[mask].to_numpy(dtype=float)
    levels_v = levels[mask].to_numpy(dtype=float)
    mins_v = mins[mask].to_numpy(dtype=float)
    n = len(ages_v)

    default_peak = DEFAULT_PEAK_AGE.get(pos, 27.0)

    if n < DEFAULT_MIN_FIT_N:
        mean_level = float(np.mean(levels_v)) if n > 0 else 50.0
        return {"a": 0.0, "b": 0.0, "c": mean_level, "peak_age": default_peak,
                "n": n, "fallback": True}

    weights = np.sqrt(np.clip(mins_v, 1.0, None))
    a, b, c = np.polyfit(ages_v, levels_v, 2, w=weights)

    if a < 0:
        peak = float(np.clip(-b / (2 * a), 18.0, 36.0))
    else:
        peak = default_peak

    return {"a": float(a), "b": float(b), "c": float(c), "peak_age": peak,
            "n": n, "fallback": False}


def fit_all_curves(df: pd.DataFrame) -> dict[str, dict]:
    curves = {}
    for pos in LEVEL_BLUEPRINTS:
        sub = df[df["_pos_group"] == pos]
        curves[pos] = fit_age_curve(sub, pos)
        tag = " (fallback — insufficient sample)" if curves[pos]["fallback"] else ""
        print(f"    {pos}: n={curves[pos]['n']:>5}  peak age ≈ {curves[pos]['peak_age']:.1f}{tag}")
    return curves


def curve_value(curve: dict, age: float) -> float:
    val = curve["a"] * age ** 2 + curve["b"] * age + curve["c"]
    return float(np.clip(val, 1.0, 99.0))


# ── Trajectory computation ──────────────────────────────────────────────────────

def trend_label(age: float, peak_age: float) -> str:
    delta = age - peak_age
    if delta <= -3:
        return "ASCENDING"
    if delta <= -1:
        return "APPROACHING PEAK"
    if delta <= 1:
        return "PEAK WINDOW"
    if delta <= 4:
        return "EARLY DECLINE"
    return "LATE DECLINE"


def dev_tier(rank: float) -> str:
    """
    Labels the 3-year outlook, not youth per se — Trajectory Score blends
    Current Level with a curve-implied 3-year projection, and because the
    fitted curves are gentle, the projection rarely swings far from today's
    level. So this tier answers 'how good will they still be in 3 years',
    which for most players tracks 'how good are they now'. Use the dedicated
    Breakout Prospects / Decline Risk sheets to isolate youth upside or
    aging risk specifically.
    """
    if rank >= 90:
        return "ELITE HORIZON"
    if rank >= 75:
        return "STRONG HORIZON"
    if rank >= 50:
        return "SOLID HORIZON"
    if rank >= 25:
        return "LIMITED HORIZON"
    return "FADING HORIZON"


def apply_trajectory(df: pd.DataFrame, curves: dict[str, dict]) -> pd.DataFrame:
    df = df.copy()
    age = pd.to_numeric(df["Age"], errors="coerce")

    predicted = pd.Series(np.nan, index=df.index)
    peak_age = pd.Series(np.nan, index=df.index)
    proj1 = pd.Series(np.nan, index=df.index)
    proj3 = pd.Series(np.nan, index=df.index)
    proj5 = pd.Series(np.nan, index=df.index)
    trend = pd.Series("", index=df.index)

    for pos, curve in curves.items():
        mask = (df["_pos_group"] == pos) & age.notna()
        if mask.sum() == 0:
            continue
        a_vals = age[mask]
        pred = a_vals.apply(lambda x: curve_value(curve, x))
        residual = df.loc[mask, "_level"] - pred

        predicted.loc[mask] = pred
        peak_age.loc[mask] = curve["peak_age"]
        proj1.loc[mask] = [curve_value(curve, x + 1) for x in a_vals] + residual
        proj3.loc[mask] = [curve_value(curve, x + 3) for x in a_vals] + residual
        proj5.loc[mask] = [curve_value(curve, x + 5) for x in a_vals] + residual
        trend.loc[mask] = a_vals.apply(lambda x: trend_label(x, curve["peak_age"]))

    df["_predicted"] = predicted
    df["_residual"] = (df["_level"] - predicted).round(2)
    df["_peak_age"] = peak_age.round(1)
    df["_years_to_peak"] = (peak_age - age).clip(lower=0).round(0)
    df["_proj1"] = proj1.clip(1, 99).round(2)
    df["_proj3"] = proj3.clip(1, 99).round(2)
    df["_proj5"] = proj5.clip(1, 99).round(2)
    df["_trend"] = trend

    df["_trajectory"] = (0.35 * df["_level"] + 0.65 * df["_proj3"]).round(2)

    df["_traj_rank"] = np.nan
    for pos in LEVEL_BLUEPRINTS:
        mask = df["_pos_group"] == pos
        if mask.sum() == 0:
            continue
        df.loc[mask, "_traj_rank"] = (df.loc[mask, "_trajectory"].rank(pct=True) * 100).round(2)

    df["_dev_tier"] = df["_traj_rank"].apply(lambda r: dev_tier(r) if pd.notna(r) else "")
    return df


# ── Build master table ──────────────────────────────────────────────────────────

OUTPUT_COLS = [
    "Player", "Team", "League", "Pos", "Full Position", "Age", "Contract",
    "Mkt Val (€)", "Minutes",
    "Current Level", "Predicted @ Age", "Vs Age Curve", "Peak Age", "Years to Peak",
    "Trend", "Proj +1Y", "Proj +3Y", "Proj +5Y", "Trajectory Score", "Trajectory Rank",
    "Dev Tier",
]


def build_master(df: pd.DataFrame) -> pd.DataFrame:
    rows: list[dict] = []
    contract_col = next((c for c in ["Contract expires", "ContractExpires"] if c in df.columns), None)
    mv_col = next((c for c in ["Market value", "MarketValue"] if c in df.columns), None)

    for _, r in df.iterrows():
        contract = r.get(contract_col, None) if contract_col else None
        if pd.notna(contract):
            try:
                contract = pd.to_datetime(contract).strftime("%Y-%m-%d")
            except Exception:
                contract = str(contract)
        else:
            contract = None

        mv = pd.to_numeric(r.get(mv_col, 0), errors="coerce") if mv_col else 0
        mv = 0 if pd.isna(mv) else mv

        rows.append({
            "Player":            r.get("Player", ""),
            "Team":              r.get("Team", ""),
            "League":            r.get("_League", ""),
            "Pos":               r.get("_pos_group", ""),
            "Full Position":     r.get("_full_position", ""),
            "Age":               int(r.get("Age", 0)) if pd.notna(r.get("Age")) else "",
            "Contract":          contract,
            "Mkt Val (€)":       int(mv),
            "Minutes":           int(r.get("_minutes", 0)),
            "Current Level":     round(float(r.get("_level", 0) or 0), 2),
            "Predicted @ Age":   round(float(r.get("_predicted", 0) or 0), 2),
            "Vs Age Curve":      float(r.get("_residual", 0) or 0),
            "Peak Age":          float(r.get("_peak_age", 0) or 0),
            "Years to Peak":     float(r.get("_years_to_peak", 0) or 0),
            "Trend":             r.get("_trend", ""),
            "Proj +1Y":          float(r.get("_proj1", 0) or 0),
            "Proj +3Y":          float(r.get("_proj3", 0) or 0),
            "Proj +5Y":          float(r.get("_proj5", 0) or 0),
            "Trajectory Score":  float(r.get("_trajectory", 0) or 0),
            "Trajectory Rank":   float(r.get("_traj_rank", 0) or 0),
            "Dev Tier":          r.get("_dev_tier", ""),
        })

    master = pd.DataFrame(rows)
    master = master.sort_values("Trajectory Score", ascending=False).reset_index(drop=True)
    return master[OUTPUT_COLS]


# ── Excel helpers ────────────────────────────────────────────────────────────────

def _fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)


def _border() -> Border:
    thin = Side(style="thin", color="CCCCCC")
    return Border(left=thin, right=thin, top=thin, bottom=thin)


def _autofit(ws) -> None:
    for col_cells in ws.columns:
        try:
            max_len = max(
                len(str(col_cells[0].value or "")),
                *(len(str(c.value or "")) for c in col_cells[1:10]),
            )
            ws.column_dimensions[col_cells[0].column_letter].width = min(max_len + 2, 40)
        except Exception:
            pass


TIER_COLORS = {
    "ELITE HORIZON":   C["breakout"],
    "STRONG HORIZON":  C["highceil"],
    "SOLID HORIZON":   C["ascending"],
    "LIMITED HORIZON": C["stable"],
    "FADING HORIZON":  C["decline"],
}
TREND_COLORS = {
    "ASCENDING":        C["trend_asc"],
    "APPROACHING PEAK": C["trend_appr"],
    "PEAK WINDOW":      C["trend_peak"],
    "EARLY DECLINE":    C["trend_early"],
    "LATE DECLINE":     C["trend_late"],
}


def write_data_sheet(ws, title: str, subtitle: str, df: pd.DataFrame) -> None:
    ws.append([title])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(len(df.columns), 10))
    ws["A1"].font = Font(bold=True, color=C["white"], size=13)
    ws["A1"].fill = _fill(C["navy"])
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 22

    ws.append([subtitle])
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max(len(df.columns), 10))
    ws["A2"].font = Font(italic=True, color=C["gold"], size=9)
    ws["A2"].fill = _fill(C["navy"])
    ws["A2"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[2].height = 16

    if df.empty:
        return

    ws.append(list(df.columns))
    hdr_row = ws.max_row
    for cell in ws[hdr_row]:
        cell.font = Font(bold=True, color=C["white"], size=9)
        cell.fill = _fill(C["header"])
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _border()
    ws.row_dimensions[hdr_row].height = 28

    tier_idx = list(df.columns).index("Dev Tier") + 1 if "Dev Tier" in df.columns else None
    trend_idx = list(df.columns).index("Trend") + 1 if "Trend" in df.columns else None
    traj_idx = list(df.columns).index("Trajectory Score") + 1 if "Trajectory Score" in df.columns else None

    for i, row_vals in enumerate(df.itertuples(index=False), start=1):
        ws.append(list(row_vals))
        data_row = ws.max_row
        bg = C["light"] if i % 2 == 0 else C["white"]
        for cell in ws[data_row]:
            cell.font = Font(size=9)
            cell.fill = _fill(bg)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = _border()

        if tier_idx:
            tc = ws.cell(data_row, tier_idx)
            hex_c = TIER_COLORS.get(str(tc.value or ""))
            if hex_c:
                tc.fill = _fill(hex_c)
                tc.font = Font(bold=True, color=C["white"], size=9)

        if trend_idx:
            tc = ws.cell(data_row, trend_idx)
            hex_c = TREND_COLORS.get(str(tc.value or ""))
            if hex_c:
                tc.fill = _fill(hex_c)
                tc.font = Font(bold=True, color=C["white"], size=9)

        if traj_idx:
            ws.cell(data_row, traj_idx).font = Font(bold=True, size=9)

    ws.freeze_panes = f"A{hdr_row + 1}"
    _autofit(ws)


def build_readme(ws, leagues: list[str], total: int, min_minutes: int) -> None:
    ws.title = "README"
    ws.sheet_view.showGridLines = False

    ws.append(["FC HRADEC KRÁLOVÉ — TRAJECTORY MODEL"])
    ws.merge_cells("A1:D1")
    ws["A1"].font = Font(bold=True, color=C["white"], size=15)
    ws["A1"].fill = _fill(C["navy"])
    ws.row_dimensions[1].height = 28

    league_label = f"{len(leagues)} leagues" if len(leagues) > 5 else " + ".join(leagues)
    ws.append([f"Waltzing Analytics  ·  Age-curve development & projection model  ·  "
               f"{league_label}  ·  {min_minutes}+ min  ·  {total:,} players"])
    ws.merge_cells("A2:D2")
    ws["A2"].font = Font(italic=True, color=C["gold"], size=10)
    ws["A2"].fill = _fill(C["navy"])
    ws.row_dimensions[2].height = 18

    ws.append([None])
    ws.append([None, "WHAT THIS MODEL DOES"])
    ws[f"B{ws.max_row}"].font = Font(bold=True, size=11, color=C["navy"])

    intro = [
        "Every other model in this workbook set (Lamberts Index, WAR, Hockey-Style) rates a",
        "player on CURRENT output. This model asks a different question: where is this player",
        "in their development curve, and what should we expect from them in 1 / 3 / 5 years?",
        "",
        "It fits a population-level 'age curve' per position from the full multi-league pool,",
        "then projects each player forward based on how far above or below that curve they",
        "currently sit. Useful for telling apart a 24-year-old who is already near their ceiling",
        "from a 24-year-old who is still climbing toward it.",
    ]
    for line in intro:
        ws.append([None, line])

    ws.append([None])
    ws.append([None, "WORKBOOK STRUCTURE"])
    ws[f"B{ws.max_row}"].font = Font(bold=True, size=11, color=C["navy"])
    ws.append([None, "Sheet", "Contents"])
    for col_letter in "BC":
        cell = ws[f"{col_letter}{ws.max_row}"]
        cell.font = Font(bold=True, color=C["white"])
        cell.fill = _fill(C["header"])
        cell.alignment = Alignment(horizontal="left")

    desc_map = {
        "README":             "This guide",
        "Age Curves":         "Fitted curve per position — peak age, coefficients, score-by-age table",
        "All Players":        f"Full {total:,}-player database ranked by Trajectory Score",
        "GK / CB / FB / DM / CM / W / FW": "Position-specific development boards",
        "Breakout Prospects": "Under-24 players furthest above the age curve — buy-low targets",
        "Decline Risk":       "High current output, already past expected peak age — aging risk / sell-high",
    }
    for sheet_name, desc in desc_map.items():
        ws.append([None, sheet_name, desc])
        row = ws.max_row
        ws[f"B{row}"].font = Font(bold=True, color=C["navy"])

    ws.append([None])
    ws.append([None, "KEY TERMS"])
    ws[f"B{ws.max_row}"].font = Font(bold=True, size=11, color=C["navy"])

    terms = [
        ("Current Level",    "Position-specific weighted percentile of current statistical output (0–100)"),
        ("Peak Age",         "Vertex of the position's fitted age curve — population's expected performance peak"),
        ("Predicted @ Age",  "What the age curve expects for a player of this age, at this position"),
        ("Vs Age Curve",     "Current Level − Predicted @ Age. Positive = ahead of schedule for their age."),
        ("Proj +1Y/+3Y/+5Y", "Age-curve value at that future age, plus the player's current residual carried forward"),
        ("Trajectory Score", "0.35 × Current Level + 0.65 × Proj +3Y — current form weighted toward near-term outlook"),
        ("Trajectory Rank",  "Percentile of Trajectory Score within the player's position group"),
        ("Dev Tier",         "3-year outlook label. Because fitted curves decline gently, this tracks current level"
                              " closely for most players — use Breakout Prospects / Decline Risk to isolate youth"
                              " upside or aging risk specifically"),
        ("ELITE HORIZON",    "Trajectory Rank ≥ 90"),
        ("STRONG HORIZON",   "Trajectory Rank ≥ 75"),
        ("SOLID HORIZON",    "Trajectory Rank ≥ 50"),
        ("LIMITED HORIZON",  "Trajectory Rank ≥ 25"),
        ("FADING HORIZON",   "Trajectory Rank < 25"),
        ("ASCENDING",        "3+ years before expected peak age"),
        ("APPROACHING PEAK", "1–3 years before expected peak age"),
        ("PEAK WINDOW",      "Within 1 year of expected peak age"),
        ("EARLY DECLINE",    "1–4 years past expected peak age"),
        ("LATE DECLINE",     "4+ years past expected peak age"),
    ]
    for term, desc in terms:
        ws.append([None, term, desc])
        ws[f"B{ws.max_row}"].font = Font(bold=True)

    ws.append([None])
    ws.append([None, "LIMITATIONS"])
    ws[f"B{ws.max_row}"].font = Font(bold=True, size=11, color=C["navy"])
    caveats = [
        "The Wyscout exports are single-season snapshots, not year-over-year player history, so",
        "the age curve is fitted CROSS-SECTIONALLY (comparing different players at different ages",
        "in the same season) rather than tracking individuals over time. This is the standard",
        "fallback when panel data isn't available, but it can understate survivorship effects",
        "(players who declined enough to leave the sampled leagues aren't in the pool). Treat",
        "projections as a population-based baseline to sanity-check judgement, not a guarantee.",
    ]
    for line in caveats:
        ws.append([None, line])

    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 85


def build_age_curve_sheet(ws, curves: dict[str, dict]) -> None:
    ws.title = "Age Curves"
    ws.sheet_view.showGridLines = False

    ws.append(["FITTED AGE CURVES — BY POSITION GROUP"])
    ws.merge_cells("A1:K1")
    ws["A1"].font = Font(bold=True, color=C["white"], size=13)
    ws["A1"].fill = _fill(C["navy"])
    ws.row_dimensions[1].height = 22

    ws.append(["Minutes-weighted quadratic fit of Current Level vs Age  ·  "
               "Peak Age = vertex of the parabola"])
    ws.merge_cells("A2:K2")
    ws["A2"].font = Font(italic=True, color=C["gold"], size=9)
    ws["A2"].fill = _fill(C["navy"])

    ws.append([None])
    ws.append(["Position", "Sample Size", "Peak Age", "Curve Fitted?", "a", "b", "c"])
    hdr = ws.max_row
    for cell in ws[hdr]:
        cell.font = Font(bold=True, color=C["white"])
        cell.fill = _fill(C["header"])
        cell.alignment = Alignment(horizontal="center")

    for pos, curve in curves.items():
        ws.append([pos, curve["n"], round(curve["peak_age"], 1),
                   "No (fallback)" if curve["fallback"] else "Yes",
                   round(curve["a"], 5), round(curve["b"], 4), round(curve["c"], 2)])
        ws.cell(ws.max_row, 1).font = Font(bold=True)

    ws.append([None])
    ws.append(["SCORE-BY-AGE REFERENCE TABLE (Current Level, 0–100)"])
    ws[f"A{ws.max_row}"].font = Font(bold=True, size=11)
    ws.append(["Age"] + list(curves.keys()))
    hdr2 = ws.max_row
    for cell in ws[hdr2]:
        cell.font = Font(bold=True, color=C["white"])
        cell.fill = _fill(C["header"])
        cell.alignment = Alignment(horizontal="center")

    for age in range(17, 39):
        row = [age] + [round(curve_value(curves[pos], age), 1) for pos in curves]
        ws.append(row)
        peak_ages = {round(c["peak_age"]) for c in curves.values()}
        if age in peak_ages:
            for col in range(1, len(row) + 1):
                ws.cell(ws.max_row, col).font = Font(bold=True)

    _autofit(ws)


def build_breakout_sheet(ws, master: pd.DataFrame) -> None:
    df = master[(master["Age"] != "") & (pd.to_numeric(master["Age"], errors="coerce") <= 23) &
                (master["Trend"] == "ASCENDING")].sort_values("Vs Age Curve", ascending=False).head(150)
    write_data_sheet(
        ws,
        f"BREAKOUT PROSPECTS — Under-24, Ascending, Ranked by Vs Age Curve",
        "Players furthest ahead of the age curve for their position while still years from peak — buy-low signal",
        df,
    )


def build_decline_sheet(ws, master: pd.DataFrame) -> None:
    """
    Flags proven, high-output players who are already past their position's
    expected peak age. The Trajectory Score gap alone rarely surfaces this —
    the fitted curves decline gently (a few points over 3 years), so this
    sheet keys off Trend + Current Level directly instead.
    """
    mask = (master["Current Level"] >= 65) & (master["Trend"].isin(["EARLY DECLINE", "LATE DECLINE"]))
    df = master[mask].copy()
    severity = {"LATE DECLINE": 0, "EARLY DECLINE": 1}
    df["_sev"] = df["Trend"].map(severity)
    df = df.sort_values(["_sev", "Current Level"], ascending=[True, False]).drop(columns="_sev").head(150)
    write_data_sheet(
        ws,
        f"DECLINE RISK — High Current Output, Past Expected Peak Age",
        "Current Level ≥ 65 and already in EARLY/LATE DECLINE for their position — proven quality, but aging risk. "
        "Short-term/loan targets or sell-high candidates rather than long-term deals",
        df,
    )


# ── Main ─────────────────────────────────────────────────────────────────────────

def run(leagues: list[str] | None, min_minutes: int, output: Path) -> None:
    print(f"\n{'='*60}")
    print("  Trajectory Model Builder")
    print(f"  Leagues: {'ALL' if leagues is None else leagues}")
    print(f"  Min minutes: {min_minutes}")
    print(f"{'='*60}\n")

    print("Loading Wyscout files…")
    raw = load_leagues(leagues, min_minutes)

    raw = add_position_group(raw)
    raw = raw[raw["_pos_group"] != "Other"].copy()
    print(f"  → {len(raw)} players with known position")

    print("Computing Current Level scores…")
    raw = compute_current_level(raw)

    print("Fitting age curves per position…")
    curves = fit_all_curves(raw)

    print("Computing trajectories and projections…")
    raw = apply_trajectory(raw, curves)

    print("Building master table…")
    master = build_master(raw)
    print(f"  → {len(master)} total players scored")

    elite_n = (master["Dev Tier"] == "ELITE HORIZON").sum()
    fading_n = (master["Dev Tier"] == "FADING HORIZON").sum()
    print(f"  → {elite_n} ELITE HORIZON  |  {fading_n} FADING HORIZON")

    print(f"\nWriting workbook → {output}")
    output.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    wb.remove(wb.active)

    print("  Writing README…")
    ws_readme = wb.create_sheet("README")
    league_list = leagues if leagues else sorted(master["League"].unique().tolist())
    build_readme(ws_readme, league_list, len(master), min_minutes)

    print("  Writing Age Curves…")
    ws_curves = wb.create_sheet("Age Curves")
    build_age_curve_sheet(ws_curves, curves)

    print("  Writing All Players…")
    ws_all = wb.create_sheet("All Players")
    write_data_sheet(
        ws_all,
        f"ALL PLAYERS — {len(master)} Ranked by Trajectory Score",
        "Full database across all positions and leagues",
        master,
    )

    pos_labels = {"GK": "GK — GOALKEEPER", "CB": "CB — CENTRE-BACK", "FB": "FB — FULL-BACK",
                  "DM": "DM — DEFENSIVE MID", "CM": "CM — CENTRAL MID", "W": "W — WINGER",
                  "FW": "FW — FORWARD"}
    for pos, label in pos_labels.items():
        print(f"  Writing {pos}…")
        grp = master[master["Pos"] == pos]
        ws_pos = wb.create_sheet(pos)
        write_data_sheet(
            ws_pos,
            f"{label} — {len(grp)} Ranked by Trajectory Score",
            f"Peak age ≈ {curves[pos]['peak_age']:.1f}  ·  Development board",
            grp,
        )

    print("  Writing Breakout Prospects…")
    ws_breakout = wb.create_sheet("Breakout Prospects")
    build_breakout_sheet(ws_breakout, master)

    print("  Writing Decline Risk…")
    ws_decline = wb.create_sheet("Decline Risk")
    build_decline_sheet(ws_decline, master)

    wb.save(output)
    print(f"\n✓ Saved {output}\n")


def main():
    parser = argparse.ArgumentParser(description="Build the Trajectory Model workbook")
    parser.add_argument("--leagues", nargs="*", default=DEFAULT_LEAGUES,
                        help="League names (Wyscout Files stems). Omit for all leagues.")
    parser.add_argument("--min-minutes", type=int, default=DEFAULT_MIN_MINUTES)
    parser.add_argument("--output", type=Path, default=OUT_DIR / "Trajectory_Model.xlsx")
    args = parser.parse_args()

    run(args.leagues, args.min_minutes, args.output)


if __name__ == "__main__":
    main()
