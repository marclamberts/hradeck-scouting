"""
build_team_performance.py
──────────────────────────
A team-performance analysis model built directly from the raw Wyscout
league exports in "Wyscout Files/", in the spirit of the style-profile /
percentile models clubs such as Brighton & Hove Albion use to read team
performance from data rather than scouting reports alone.

What it does
------------
1. Loads every league export in "Wyscout Files/" and tags each player row
   with League / Country / Tier (via the Leagues Overview tier system).
2. Rolls players up to their club ("Team within selected timeframe" — the
   club they actually featured for during the sampled window, which is a
   cleaner join key than "Team" for players who moved/loaned mid-season).
3. Builds two kinds of team numbers:
     - Literal output: goals/xG for & against per 90 (goals-against and
       xG-against are read off the team's own goalkeeper rows, which is
       the one genuinely team-level signal hiding inside a player export).
     - Style pillars: minutes-weighted per-90 rates across Attacking,
       Creation, Buildup, Directness, Pressing/Defending, Physical and
       Discipline, each converted to a percentile rank *within tier* so
       teams are compared against peers of similar competitive strength.
4. Combines the pillars + net output percentile into a single
   Team Performance Index (0-100), tier-weighted for a cross-league
   Overall Rank, mirroring League Analysis/Team Ratings.xlsx.
5. Writes League Analysis/Team Performance.xlsx (Methodology, Overall
   Rank, per-tier and per-league sheets).
6. Optionally renders a Brighton-style percentile pizza chart for one
   club: `--radar "Team Name"`.

Usage
-----
  python build_team_performance.py
  python build_team_performance.py --min-minutes 3000 --min-players 5
  python build_team_performance.py --radar "Brighton" --league England
"""
from __future__ import annotations

import argparse
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ── Paths ────────────────────────────────────────────────────────────────────
ROOT = Path(__file__).parent
WYSCOUT_DIR = ROOT / "Wyscout Files"
OUT_DIR = ROOT / "League Analysis"
OUT_DIR.mkdir(exist_ok=True)
OUT_XLSX = OUT_DIR / "Team Performance.xlsx"
REPORTS_DIR = ROOT / "reports"
REPORTS_DIR.mkdir(exist_ok=True)

DEFAULT_MIN_TEAM_MINUTES = 3000
DEFAULT_MIN_PLAYERS = 5

TEAM_COL = "Team within selected timeframe"

# ── Wyscout stem → (Country label for join, Division number) ─────────────────
# Multi-part files collapse onto the same (country, division).
STEM_COUNTRY_DIV: dict[str, tuple[str, int]] = {
    "England": ("England", 1), "Spain": ("Spain", 1), "Germany": ("Germany", 1),
    "Italy": ("Italy", 1), "France": ("France", 1),
    "Argentina": ("Argentina", 1), "Brazil": ("Brazil", 1), "Mexico": ("Mexico", 1),
    "Netherlands": ("Netherlands", 1), "Belgium": ("Belgium", 1), "Turkiye": ("Türkiye", 1),
    "Portugal": ("Portugal", 1), "Russia": ("Russia", 1), "Ukraine": ("Ukraine", 1),
    "Scotland": ("Scotland", 1), "Japan": ("Japan", 1), "USA": ("USA", 1),
    "Korea": ("South Korea", 1), "Saudi": ("Saudi Arabia", 1), "China": ("China", 1),
    "Qatar": ("Qatar", 1), "UAE": ("UAE", 1),
    "England II": ("England", 2), "Spain II": ("Spain", 2), "Germany II": ("Germany", 2),
    "Italy II": ("Italy", 2), "France II": ("France", 2),
    "Austria": ("Austria", 1), "Switzerland": ("Switzerland", 1), "Denmark": ("Denmark", 1),
    "Sweden": ("Sweden", 1), "Norway": ("Norway", 1), "Poland": ("Poland", 1),
    "Czech": ("Czech Republic", 1), "Greece": ("Greece", 1), "Croatia": ("Croatia", 1),
    "Romania": ("Romania", 1), "Serbia": ("Serbia", 1), "Colombia": ("Colombia", 1),
    "Chile": ("Chile", 1), "Peru": ("Peru", 1), "Uruguay": ("Uruguay", 1),
    "Australia": ("Australia", 1), "Egypt": ("Egypt", 1), "Morocco": ("Morocco", 1),
    "Nigeria": ("Nigeria", 1), "South Africa": ("South Africa", 1), "India": ("India", 1),
    "Indonesia": ("Indonesia", 1), "Malaysia": ("Malaysia", 1), "Thailand": ("Thailand", 1),
    "Canada": ("Canada", 1),
    "England III": ("England", 3), "Spain III": ("Spain", 3), "Germany III": ("Germany", 3),
    "France III": ("France", 3),
    "Austria II": ("Austria", 2), "Switzerland II": ("Switzerland", 2), "Denmark II": ("Denmark", 2),
    "Sweden II": ("Sweden", 2), "Norway II": ("Norway", 2), "Poland II": ("Poland", 2),
    "Czech II": ("Czech Republic", 2), "Greece II": ("Greece", 2), "Russia II": ("Russia", 2),
    "Ukraine II": ("Ukraine", 2), "Scotland II": ("Scotland", 2), "Japan II III": ("Japan", 2),
    "USA II": ("USA", 2), "Korea II": ("South Korea", 2), "Saudi II": ("Saudi Arabia", 2),
    "China II": ("China", 2), "Portugal II": ("Portugal", 2), "Belgium II": ("Belgium", 2),
    "Turkiye II": ("Türkiye", 2), "Netherlands II": ("Netherlands", 2), "Serbia II": ("Serbia", 2),
    "Argentina II": ("Argentina", 2), "Brazil II": ("Brazil", 2), "Mexico II": ("Mexico", 2),
    "Chile II": ("Chile", 2), "Uruguay II": ("Uruguay", 2),
    "Ecuador": ("Ecuador", 1), "Ecuador II": ("Ecuador", 2),
    "Albania": ("Albania", 1), "Andorra": ("Andorra", 1), "Armenia": ("Armenia", 1),
    "Azerbaijan": ("Azerbaijan", 1), "Bahrain": ("Bahrain", 1), "Bangladesh": ("Bangladesh", 1),
    "Bolivia": ("Bolivia", 1), "Bosnia": ("Bosnia", 1), "Bulgaria": ("Bulgaria", 1),
    "Cambodia": ("Cambodia", 1), "Costa Rica": ("Costa Rica", 1), "Cyprus": ("Cyprus", 1),
    "Finland": ("Finland", 1), "Georgia": ("Georgia", 1), "Guatemala": ("Guatemala", 1),
    "Honduras": ("Honduras", 1), "Hong Kong": ("Hong Kong", 1), "Hungary": ("Hungary", 1),
    "Iceland": ("Iceland", 1), "Ireland": ("Ireland", 1), "Jordan": ("Jordan", 1),
    "Kazakhstan": ("Kazakhstan", 1), "Kosovo": ("Kosovo", 1), "Kyrgystan": ("Kyrgyzstan", 1),
    "Latvia": ("Latvia", 1), "Lithuania": ("Lithuania", 1), "Malta": ("Malta", 1),
    "Moldovia": ("Moldova", 1), "Montenegro": ("Montenegro", 1), "Nicaragua": ("Nicaragua", 1),
    "Northern Ireland": ("Northern Ireland", 1), "Panama": ("Panama", 1), "Paraguay": ("Paraguay", 1),
    "Philippines": ("Philippines", 1), "Singapore": ("Singapore", 1), "Slovakia": ("Slovakia", 1),
    "Slovenia": ("Slovenia", 1), "Tunisia": ("Tunisia", 1), "Vietnam": ("Vietnam", 1),
    "Wales": ("Wales", 1), "Uzbekistan": ("Uzbekistan", 1), "Venezuela": ("Venezuela", 1),
    "El Salvador": ("El Salvador", 1), "Estonia": ("Estonia", 1), "Faroe Islands": ("Faroe Islands", 1),
    "England IV": ("England", 4), "England V": ("England", 5),
    "Norway III": ("Norway", 3), "Denmark III": ("Denmark", 3), "Denmark IV": ("Denmark", 4),
    "Poland III": ("Poland", 3), "Portugal III": ("Portugal", 3), "Netherlands III": ("Netherlands", 3),
    "Sweden III": ("Sweden", 3), "Korea III": ("South Korea", 3), "Scotland III": ("Scotland", 3),
    "Scotland IV": ("Scotland", 4), "Cyprus II": ("Cyprus", 2), "Hungary II": ("Hungary", 2),
    "Ireland II": ("Ireland", 2), "Slovakia II": ("Slovakia", 2), "Slovenia II": ("Slovenia", 2),
    "Finland II": ("Finland", 2), "USA III": ("USA", 3),
    "Italy III - Part I": ("Italy", 3), "Italy III - Part II": ("Italy", 3),
    "Italy III - Part III": ("Italy", 3), "Italy III - Part IV": ("Italy", 3),
    "Germany 4 - Part I": ("Germany", 4), "Germany 4 - Part II": ("Germany", 4),
    "Germany 4 - Part III": ("Germany", 4), "Germany 4 - Part IV": ("Germany", 4),
    "Australia II - Part I": ("Australia", 2), "Australia II - Part II": ("Australia", 2),
    "Australia II - Part III": ("Australia", 2), "Australia II - Part IV": ("Australia", 2),
    "Australia II - Part V": ("Australia", 2), "Australia II - Part VI": ("Australia", 2),
    "Australia II - Part VII": ("Australia", 2),
    "Czech U17": ("Czech Republic", 99), "Czech U19": ("Czech Republic", 99),
}

TIER_WEIGHT = {1: 1.00, 2: 0.93, 3: 0.86, 4: 0.79, 5: 0.72, 6: 0.65}

# ── Style pillars (all "per 90", minutes-weighted across the squad) ──────────
ATTACK_METRICS = [
    "Goals per 90", "xG per 90", "Shots per 90", "Shots on target, %",
    "Touches in box per 90", "Goal conversion, %",
]
CREATION_METRICS = [
    "xA per 90", "Key passes per 90", "Smart passes per 90", "Through passes per 90",
    "Deep completions per 90", "Passes to penalty area per 90", "Crosses per 90",
]
BUILDUP_METRICS = [
    "Passes per 90", "Accurate passes, %", "Progressive passes per 90",
    "Passes to final third per 90", "Accurate progressive passes, %",
]
DIRECTNESS_METRICS = [
    "Average pass length, m", "Long passes per 90", "Accurate long passes, %",
    "Progressive runs per 90", "Accelerations per 90",
]
PRESSING_METRICS = [
    "Successful defensive actions per 90", "PAdj Interceptions", "PAdj Sliding tackles",
    "Defensive duels per 90", "Defensive duels won, %", "Shots blocked per 90",
]
PHYSICAL_METRICS = [
    "Duels per 90", "Duels won, %", "Aerial duels per 90", "Aerial duels won, %",
]
DISCIPLINE_METRICS = ["Fouls per 90", "Yellow cards per 90", "Red cards per 90"]  # lower is better

PILLARS: dict[str, list[str]] = {
    "AttackingScore": ATTACK_METRICS,
    "CreationScore": CREATION_METRICS,
    "BuildupScore": BUILDUP_METRICS,
    "PressingScore": PRESSING_METRICS,
    "PhysicalScore": PHYSICAL_METRICS,
    "DisciplineScore": DISCIPLINE_METRICS,
}
INVERT_PILLARS = {"DisciplineScore"}

# Style axis reported for context but NOT part of the composite index —
# directness is a playing-style choice, not a quality signal.
STYLE_ONLY_PILLARS = {"DirectnessScore": DIRECTNESS_METRICS}

INDEX_WEIGHTS = {
    "AttackingScore": 0.18, "CreationScore": 0.14, "BuildupScore": 0.10,
    "PressingScore": 0.18, "PhysicalScore": 0.08, "DisciplineScore": 0.05,
    "NetOutputScore": 0.27,
}

RADAR_METRICS = [
    ("Goals per 90", "Goals"), ("xG per 90", "xG"), ("Shots per 90", "Shots"),
    ("Touches in box per 90", "Box touches"),
    ("xA per 90", "xA"), ("Key passes per 90", "Key passes"), ("Crosses per 90", "Crosses"),
    ("Accurate passes, %", "Pass acc."), ("Progressive passes per 90", "Prog. passes"),
    ("Successful defensive actions per 90", "Def. actions"), ("PAdj Interceptions", "Interceptions"),
    ("Defensive duels won, %", "Duels won"), ("Aerial duels won, %", "Aerial won"),
]
RADAR_SLICE_COLOURS = (
    ["#E4572E"] * 4 + ["#2E86AB"] * 5 + ["#2E7D32"] * 4
)


# ══════════════════════════════════════════════════════════════════════════════
# Load & tag every league file
# ══════════════════════════════════════════════════════════════════════════════
def _tier_lookup() -> dict[tuple[str, int], tuple[int, str]]:
    leagues_ov = pd.read_excel(ROOT / "data" / "Leagues Overview.xlsx")
    tier_map: dict[tuple[str, int], tuple[int, str]] = {}
    for _, row in leagues_ov.iterrows():
        div = row["Division"]
        if isinstance(div, int) or (isinstance(div, str) and str(div).isdigit()):
            tier_map[(str(row["Country"]).strip(), int(div))] = (
                int(row["Tier"]), str(row["Tier Label"])
            )
    return tier_map


def load_all_players() -> pd.DataFrame:
    tier_map = _tier_lookup()

    def lookup_tier(stem: str) -> tuple[int, str]:
        if stem not in STEM_COUNTRY_DIV:
            return (4, "Developing")
        country, div = STEM_COUNTRY_DIV[stem]
        if div == 99:
            return (6, "Youth/Grassroots")
        return tier_map.get((country, div), (4, "Developing"))

    files = sorted(WYSCOUT_DIR.glob("*.xlsx"))
    print(f"Loading {len(files)} Wyscout league files …")
    frames = []
    for path in files:
        tier, tier_label = lookup_tier(path.stem)
        if tier == 6:
            continue  # youth leagues excluded from team-performance model
        try:
            df = pd.read_excel(path)
        except Exception as exc:
            print(f"  [skip] {path.stem}: {exc}")
            continue
        df.columns = [str(c).strip() for c in df.columns]
        if TEAM_COL not in df.columns:
            print(f"  [skip] {path.stem}: no '{TEAM_COL}' column")
            continue
        df = df.assign(League=path.stem, Country=STEM_COUNTRY_DIV.get(path.stem, ("?", 0))[0],
                        Tier=tier, TierLabel=tier_label)
        frames.append(df)

    raw = pd.concat(frames, ignore_index=True)

    meta_cols = {"Player", "Team", TEAM_COL, "Position", "Birth country", "Passport country",
                 "Foot", "On loan", "Contract expires", "League", "Country", "TierLabel"}
    for col in raw.columns:
        if col not in meta_cols:
            raw[col] = pd.to_numeric(raw[col], errors="coerce")

    raw = raw.dropna(subset=[TEAM_COL])
    raw = raw[raw[TEAM_COL].astype(str).str.strip() != ""]
    raw["Minutes played"] = raw["Minutes played"].fillna(0)
    raw["_is_gk"] = raw["Position"].astype(str).str.contains("GK")
    print(f"  → {len(raw):,} player rows across {raw['League'].nunique()} leagues")
    return raw.reset_index(drop=True)


# ══════════════════════════════════════════════════════════════════════════════
# Roll players up to clubs
# ══════════════════════════════════════════════════════════════════════════════
def _weighted_mean(df: pd.DataFrame, col: str, weight_col: str = "Minutes played") -> float:
    w = df[weight_col].clip(lower=0)
    vals = df[col]
    mask = vals.notna() & (w > 0)
    if not mask.any():
        return np.nan
    return float((vals[mask] * w[mask]).sum() / w[mask].sum())


def build_team_table(min_minutes: int, min_players: int) -> pd.DataFrame:
    rows = []
    group_keys = ["League", "Country", "Tier", "TierLabel", TEAM_COL]
    for (league, country, tier, tier_label, team), g in raw_df.groupby(group_keys):
        total_minutes = g["Minutes played"].sum()
        n_players = len(g)
        if total_minutes < min_minutes or n_players < min_players:
            continue

        rec = {
            "Team": team, "League": league, "Country": country,
            "Tier": tier, "TierLabel": tier_label,
            "Players": n_players, "TotalMinutes": int(total_minutes),
        }

        for pillar_metrics in (ATTACK_METRICS, CREATION_METRICS, BUILDUP_METRICS,
                               DIRECTNESS_METRICS, PRESSING_METRICS, PHYSICAL_METRICS,
                               DISCIPLINE_METRICS):
            for m in pillar_metrics:
                if m in g.columns:
                    rec[m] = _weighted_mean(g, m)

        # Literal net-output record, sourced from the club's own goalkeeper(s)
        gk = g[g["_is_gk"]]
        gk_minutes = gk["Minutes played"].sum()
        total_goals = g["Goals"].sum() if "Goals" in g.columns else np.nan
        total_xg = g["xG"].sum() if "xG" in g.columns else np.nan
        if gk_minutes > 0:
            conceded = gk["Conceded goals"].sum() if "Conceded goals" in gk.columns else np.nan
            total_xg_against = gk["xG against"].sum() if "xG against" in gk.columns else np.nan
            rec["TeamMatchMinutes"] = int(gk_minutes)
            rec["GoalsForPer90"] = 90 * total_goals / gk_minutes
            rec["GoalsAgainstPer90"] = 90 * conceded / gk_minutes
            rec["xGForPer90"] = 90 * total_xg / gk_minutes
            rec["xGAgainstPer90"] = 90 * total_xg_against / gk_minutes
            rec["GoalDifferencePer90"] = rec["GoalsForPer90"] - rec["GoalsAgainstPer90"]
            rec["xGDifferencePer90"] = rec["xGForPer90"] - rec["xGAgainstPer90"]
        else:
            for c in ("TeamMatchMinutes", "GoalsForPer90", "GoalsAgainstPer90",
                      "xGForPer90", "xGAgainstPer90", "GoalDifferencePer90", "xGDifferencePer90"):
                rec[c] = np.nan

        rows.append(rec)

    return pd.DataFrame(rows)


# ══════════════════════════════════════════════════════════════════════════════
# Percentile pillars + composite index
# ══════════════════════════════════════════════════════════════════════════════
def add_scores(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()

    def tier_pct(col: str, ascending: bool = True) -> pd.Series:
        return df.groupby("Tier")[col].rank(pct=True, ascending=ascending) * 100

    all_pillars = {**PILLARS, **STYLE_ONLY_PILLARS}
    for pillar, metrics in all_pillars.items():
        present = [m for m in metrics if m in df.columns]
        if not present:
            df[pillar] = np.nan
            continue
        ascending = pillar not in INVERT_PILLARS
        pct_cols = []
        for m in present:
            pc = f"_pct_{pillar}_{m}"
            df[pc] = tier_pct(m, ascending=ascending)
            pct_cols.append(pc)
        df[pillar] = df[pct_cols].mean(axis=1)
        df.drop(columns=pct_cols, inplace=True)

    # Net output percentile (within tier), from literal xG difference
    df["NetOutputScore"] = df.groupby("Tier")["xGDifferencePer90"].rank(pct=True) * 100

    weighted = sum(df[k].fillna(df[k].mean()) * w for k, w in INDEX_WEIGHTS.items())
    df["TeamPerformanceIndex"] = weighted.round(1)

    df["TierWeight"] = df["Tier"].map(TIER_WEIGHT).fillna(0.65)
    df["OverallIndex"] = (df["TeamPerformanceIndex"] * df["TierWeight"]).round(1)

    df["LeagueRank"] = df.groupby("League")["TeamPerformanceIndex"].rank(
        method="min", ascending=False).astype(int)
    df["TierRank"] = df.groupby("Tier")["TeamPerformanceIndex"].rank(
        method="min", ascending=False).astype(int)
    df["OverallRank"] = df["OverallIndex"].rank(method="min", ascending=False).astype(int)

    round_cols = [c for c in all_pillars] + [
        "NetOutputScore", "GoalsForPer90", "GoalsAgainstPer90", "xGForPer90",
        "xGAgainstPer90", "GoalDifferencePer90", "xGDifferencePer90",
    ]
    for c in round_cols:
        if c in df.columns:
            df[c] = df[c].round(2)

    return df


# ══════════════════════════════════════════════════════════════════════════════
# Excel styling (mirrors League Analysis/Team Ratings.xlsx house style)
# ══════════════════════════════════════════════════════════════════════════════
TIER_COLOURS = {1: "1A3A6B", 2: "1565C0", 3: "2E7D32", 4: "E65100", 5: "6A1E55", 6: "4E342E"}
TIER_LIGHT = {1: "DDEEFF", 2: "E3F2FD", 3: "E8F5E9", 4: "FFF3E0", 5: "F3E5F5", 6: "EFEBE9"}
HEADER_BG, WHITE = "2C3E50", "FFFFFF"

DISPLAY_COLS = [
    "OverallRank", "TierRank", "LeagueRank", "Team", "League", "Country", "TierLabel",
    "Players", "TotalMinutes",
    "TeamPerformanceIndex", "OverallIndex",
    "GoalsForPer90", "GoalsAgainstPer90", "GoalDifferencePer90",
    "xGForPer90", "xGAgainstPer90", "xGDifferencePer90", "NetOutputScore",
    "AttackingScore", "CreationScore", "BuildupScore", "DirectnessScore",
    "PressingScore", "PhysicalScore", "DisciplineScore",
]


def _thin_border() -> Border:
    s = Side(style="thin", color="CCCCCC")
    return Border(left=s, right=s, top=s, bottom=s)


def _header_style(ws, row: int, n_cols: int, bg: str = HEADER_BG) -> None:
    for c in range(1, n_cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = PatternFill("solid", fgColor=bg)
        cell.font = Font(bold=True, color=WHITE, size=10)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _thin_border()


def _write_sheet(writer, sheet_name: str, df: pd.DataFrame, tier_num: int | None = None) -> None:
    out = df[DISPLAY_COLS].copy()
    out.to_excel(writer, sheet_name=sheet_name[:31], index=False)
    ws = writer.sheets[sheet_name[:31]]
    ws.row_dimensions[1].height = 34
    _header_style(ws, 1, len(out.columns))

    for r in range(2, len(out) + 2):
        tval = ws.cell(row=r, column=7).value  # TierLabel column
        try:
            tnum = int(str(tval).split()[-1])
        except Exception:
            tnum = tier_num or 1
        bg = TIER_LIGHT.get(tnum, WHITE) if r % 2 == 0 else WHITE
        for c in range(1, len(out.columns) + 1):
            cell = ws.cell(row=r, column=c)
            cell.fill = PatternFill("solid", fgColor=bg)
            cell.border = _thin_border()
            if c <= 3:
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.font = Font(bold=True, size=9, color=TIER_COLOURS.get(tnum, "000000"))
            elif c in (4, 5, 6, 7):
                cell.alignment = Alignment(horizontal="left", vertical="center")
                cell.font = Font(size=9, bold=(c == 4))
            else:
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.font = Font(size=9)

    widths = [10, 9, 9, 24, 26, 16, 12, 8, 12, 16, 12, 12, 14, 14, 10, 12, 12, 12,
              12, 12, 11, 12, 11, 11, 12]
    for i, w in enumerate(widths[: len(out.columns)], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"


def _write_methodology(writer) -> None:
    lines = [
        ("Team Performance Model — Methodology", True),
        ("", False),
        ("Source data", True),
        ("Every workbook in \"Wyscout Files/\" (one per league) is a Wyscout player search "
         "export. There is no native team-level table, so this model rebuilds one by rolling "
         "players up to \"Team within selected timeframe\" — the club a player actually turned "
         "out for in the sampled window.", False),
        ("", False),
        ("Two kinds of number", True),
        ("1. Net output (literal): Goals/xG For and Against per 90, read from the summed "
         "goal/xG involvement of every player on the roster and the club's own goalkeeper "
         "rows (Conceded goals, xG against). This is the closest thing to a real team result "
         "hiding inside a player export.", False),
        ("2. Style pillars (relative): minutes-weighted per-90 rates across Attacking, "
         "Creation, Buildup, Pressing/Defending, Physical duels and Discipline, each converted "
         "to a percentile rank within the team's own tier so a Championship team is judged "
         "against Championship peers, not Premier League ones. Directness is reported as a "
         "style axis only — it describes HOW a team plays, not how well.", False),
        ("", False),
        ("Team Performance Index", True),
        ("Weighted blend of the six style pillars (73%) and the Net Output percentile (27%), "
         "0-100 within tier. OverallIndex multiplies this by a tier weight (Elite 1.00 down to "
         "Lower 0.72) for a single cross-league Overall Rank, the same convention used in "
         "Team Ratings.xlsx.", False),
        ("", False),
        ("Caveats", True),
        ("Wyscout Files/ caps most leagues at ~500 player rows, so fringe squad players can be "
         "under-represented; teams below the minutes/player-count threshold are dropped rather "
         "than shown on partial data. Net-output figures assume the sampled goalkeeper(s) "
         "covered most of the club's minutes.", False),
    ]
    ws = writer.book.create_sheet("Methodology", 0)
    ws.column_dimensions["A"].width = 110
    for i, (text, bold) in enumerate(lines, 1):
        cell = ws.cell(row=i, column=1, value=text)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        cell.font = Font(bold=bold, size=12 if bold and i == 1 else (10.5 if bold else 10),
                          color=HEADER_BG if bold else "000000")
        ws.row_dimensions[i].height = 18 if not text else (28 if bold else 32)


def write_workbook(df: pd.DataFrame) -> None:
    with pd.ExcelWriter(OUT_XLSX, engine="openpyxl") as writer:
        _write_methodology(writer)

        _write_sheet(writer, "All Teams (Overall Rank)", df.sort_values("OverallRank"))

        for tier_num in sorted(df["Tier"].unique()):
            sub = df[df["Tier"] == tier_num].sort_values("TierRank")
            if sub.empty:
                continue
            label = sub["TierLabel"].iloc[0]
            _write_sheet(writer, f"Tier {tier_num} - {label}", sub, tier_num=int(tier_num))

        for league in sorted(df["League"].unique()):
            sub = df[df["League"] == league].sort_values("LeagueRank")
            tier_num = int(sub["Tier"].iloc[0])
            _write_sheet(writer, league, sub, tier_num=tier_num)

    print(f"Saved: {OUT_XLSX}")


# ══════════════════════════════════════════════════════════════════════════════
# Brighton-style percentile pizza chart for a single club
# ══════════════════════════════════════════════════════════════════════════════
def render_radar(df: pd.DataFrame, raw_df: pd.DataFrame, team: str, league: str | None = None) -> Path:
    from mplsoccer import PyPizza
    import matplotlib.pyplot as plt

    sub = df[df["Team"].str.lower() == team.lower()]
    if league:
        sub = sub[sub["League"].str.lower() == league.lower()]
    if sub.empty:
        raise SystemExit(f"No team-performance row found for '{team}'"
                          + (f" in league '{league}'" if league else "")
                          + ". Check spelling / minutes threshold.")
    if len(sub) > 1:
        opts = ", ".join(f"{r.Team} ({r.League})" for r in sub.itertuples())
        raise SystemExit(f"Multiple matches for '{team}': {opts}. Pass --league to disambiguate.")
    row = sub.iloc[0]

    players = raw_df[(raw_df[TEAM_COL] == row["Team"]) & (raw_df["League"] == row["League"])]
    tier_peers = raw_df[raw_df["Tier"] == row["Tier"]]

    values, params = [], []
    for col, label in RADAR_METRICS:
        team_val = _weighted_mean(players, col)
        peer_vals = tier_peers.groupby(TEAM_COL).apply(lambda g: _weighted_mean(g, col))
        peer_vals = peer_vals.dropna()
        pct = (peer_vals < team_val).mean() * 100 if pd.notna(team_val) and len(peer_vals) else 0
        values.append(round(pct))
        params.append(label)

    baker = PyPizza(
        params=params, min_range=[0] * len(params), max_range=[100] * len(params),
        background_color="#0E1117", straight_line_color="#39424E", straight_line_lw=1,
        last_circle_color="#39424E", last_circle_lw=2, other_circle_lw=0,
        inner_circle_size=6,
    )
    fig, ax = baker.make_pizza(
        values, figsize=(9, 9.5),
        color_blank_space="same",
        slice_colors=RADAR_SLICE_COLOURS,
        value_colors=["#0E1117"] * len(params),
        value_bck_colors=RADAR_SLICE_COLOURS,
        kwargs_slices=dict(edgecolor="#0E1117", zorder=2, linewidth=1.5),
        kwargs_params=dict(color="#E8E8E8", fontsize=10, va="center"),
        kwargs_values=dict(
            color="#0E1117", fontsize=10, zorder=3,
            bbox=dict(edgecolor="#0E1117", facecolor="cornflowerblue",
                      boxstyle="round,pad=0.2", lw=1),
        ),
    )
    fig.text(0.5, 0.975, row["Team"], size=20, ha="center", color="#FFFFFF", weight="bold")
    fig.text(0.5, 0.945,
              f"Team style percentile vs {row['TierLabel']} peers  |  {row['League']}  |  "
              f"{int(row['TotalMinutes']):,} squad minutes sampled",
              size=10.5, ha="center", color="#B0B6BE")
    fig.text(0.03, 0.02,
              "Red = Attacking   Blue = Creation & Buildup   Green = Pressing & Duels",
              size=8.5, color="#B0B6BE")

    out_path = REPORTS_DIR / f"Team_Style_{row['Team'].replace(' ', '_')}.png"
    fig.savefig(out_path, dpi=200, facecolor="#0E1117", bbox_inches="tight")
    plt.close(fig)
    print(f"Saved radar: {out_path}")
    return out_path


# ══════════════════════════════════════════════════════════════════════════════
raw_df: pd.DataFrame | None = None


def main() -> None:
    global raw_df
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--min-minutes", type=int, default=DEFAULT_MIN_TEAM_MINUTES,
                     help="Minimum sampled squad minutes for a team to be included")
    ap.add_argument("--min-players", type=int, default=DEFAULT_MIN_PLAYERS,
                     help="Minimum sampled players for a team to be included")
    ap.add_argument("--radar", type=str, default=None, help="Render a percentile pizza chart for one club")
    ap.add_argument("--league", type=str, default=None, help="Disambiguate --radar by league/file stem")
    args = ap.parse_args()

    raw_df = load_all_players()

    print("Rolling players up to clubs …")
    team_df = build_team_table(args.min_minutes, args.min_players)
    print(f"  → {len(team_df):,} clubs kept "
          f"(>= {args.min_minutes} minutes, >= {args.min_players} players)")

    print("Scoring style pillars + net output …")
    team_df = add_scores(team_df)

    print("Writing workbook …")
    write_workbook(team_df)

    if args.radar:
        render_radar(team_df, raw_df, args.radar, args.league)


if __name__ == "__main__":
    main()
