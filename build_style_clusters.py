"""
build_style_clusters.py
────────────────────────
Style Cluster Model — data-driven playing-style archetypes.

Every other model in this repo (Lamberts Index, WAR, Hockey-Style, Wyscout
composite scores, Trajectory Model) rates players on a QUALITY axis — how
good is this player. This model asks a different question entirely: what
KIND of player is this, regardless of how good they are.

Methodology
-----------
1. Style metrics only — per position group, a curated set of per-90 rate /
   tendency metrics that describe role and process (passing tendencies,
   crossing, carrying, aerial involvement, pressing volume, zone of
   operation). Pure output/quality metrics (Goals, xG, Assists, xA, Goal
   conversion %) are deliberately EXCLUDED — including them would just
   split players into "good" vs "bad" again, not "how they play".

2. Standardize — z-score each style metric within its position group so no
   single metric's scale dominates.

3. Unsupervised clustering — KMeans on the standardized style vectors.
   The number of clusters k is chosen per position by silhouette score over
   a small candidate range (not hand-picked), so the archetypes that emerge
   reflect actual structure in the data rather than a preset taxonomy.

4. Auto-labelling — each cluster's label is generated directly from its
   centroid: the two style metrics with the highest positive z-score in
   that cluster become the label (e.g. "High Crosses + High Progressive
   Runs"). No hand-curated archetype names — the label is a direct,
   falsifiable readout of the data, not house style.

5. Style Fit — a player's percentile distance-to-centroid within their
   assigned cluster. 100 = the most prototypical example of that style in
   the pool; low values mean they're a borderline / hybrid case.

6. Prototype players — for each cluster, the 3 players closest to the
   centroid, so a scout can anchor an unfamiliar cluster label to players
   they already know.

Output: reports/Style_Cluster_Model.xlsx
  README · Style Clusters Overview · GK/CB/FB/DM/CM/W/FW

Usage:
  python build_style_clusters.py
  python build_style_clusters.py --leagues "Czech II" Slovakia --min-minutes 700
  python build_style_clusters.py --output reports/My_Style_Model.xlsx
"""
from __future__ import annotations

import argparse
import re
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from sklearn.cluster import KMeans
from sklearn.metrics import silhouette_score
from sklearn.preprocessing import StandardScaler

# ── Config ─────────────────────────────────────────────────────────────────────

ROOT = Path(__file__).parent
WYSCOUT_DIR = ROOT / "Wyscout Files"
OUT_DIR = ROOT / "reports"

SKIP_FILES = {"FCHK Model V3 - Loaded Leagues", "FCHK Model V3 - Model Input",
              "FCHK Model V3 - Player Scores", "FCHK Model V3 - Player Styles",
              "FCHK Model V3 - Recruitment Scores", "FCHK Model V3 - Smart Club Closeness",
              "FCHK Model V3 - Summary", "FCHK Model V3 Scores", "FCHK Scouting Report",
              "Leagues Overview", "Wyscout Anomaly Report", "Wyscout Full Scouting Report"}

DEFAULT_LEAGUES = None  # None = all leagues in WYSCOUT_DIR
DEFAULT_MIN_MINUTES = 600
K_MIN, K_MAX = 3, 6
MIN_SAMPLES_PER_K = 15   # require at least this many players per candidate cluster
MIN_POOL_SIZE = 30       # below this, skip clustering for that position entirely

POS_MAP: dict[str, str] = {
    "CF": "FW", "SS": "FW",
    "LW": "W",  "RW": "W",  "LWF": "W",  "RWF": "W",  "WF": "W",
    "LAMF": "W", "RAMF": "W",
    "AMF": "CM",
    "CMF": "CM", "LCM": "CM", "RCM": "CM", "LCMF": "CM", "RCMF": "CM",
    "DMF": "DM", "LDM": "DM", "RDM": "DM", "LDMF": "DM", "RDMF": "DM",
    "LB": "FB",  "RB": "FB",  "LWB": "FB", "RWB": "FB",
    "CB": "CB",  "LCB": "CB", "RCB": "CB",
    "GK": "GK",
}

# Style metrics — role/process indicators only, no output/quality metrics
STYLE_METRICS: dict[str, list[str]] = {
    "GK": [
        "Exits per 90", "Aerial duels per 90.1", "Back passes received as GK per 90",
        "Passes per 90", "Accurate passes, %", "Long passes per 90",
        "Accurate long passes, %", "Average pass length, m",
    ],
    "CB": [
        "Aerial duels per 90", "Aerial duels won, %", "Sliding tackles per 90",
        "Interceptions per 90", "Defensive duels per 90", "Progressive passes per 90",
        "Progressive runs per 90", "Long passes per 90", "Accurate long passes, %",
        "Fouls per 90",
    ],
    "FB": [
        "Crosses per 90", "Accurate crosses, %", "Progressive runs per 90",
        "Dribbles per 90", "Successful dribbles, %", "Defensive duels per 90",
        "Aerial duels per 90", "Accelerations per 90", "Progressive passes per 90",
        "Deep completed crosses per 90",
    ],
    "DM": [
        "Passes per 90", "Accurate passes, %", "Progressive passes per 90",
        "Interceptions per 90", "Defensive duels per 90", "Aerial duels per 90",
        "Long passes per 90", "Fouls per 90", "Through passes per 90",
        "Short / medium passes per 90",
    ],
    "CM": [
        "Passes per 90", "Progressive passes per 90", "Key passes per 90",
        "Through passes per 90", "Smart passes per 90", "Progressive runs per 90",
        "Dribbles per 90", "Defensive duels per 90", "Long passes per 90",
        "Deep completions per 90",
    ],
    "W": [
        "Dribbles per 90", "Successful dribbles, %", "Crosses per 90",
        "Accurate crosses, %", "Progressive runs per 90", "Touches in box per 90",
        "Accelerations per 90", "Key passes per 90", "Through passes per 90",
        "Offensive duels per 90",
    ],
    "FW": [
        "Touches in box per 90", "Aerial duels per 90", "Aerial duels won, %",
        "Dribbles per 90", "Offensive duels per 90", "Shots per 90",
        "Received long passes per 90", "Progressive runs per 90",
        "Key passes per 90", "Passes per 90",
    ],
}

# ── Colors ─────────────────────────────────────────────────────────────────────
C = {
    "navy":   "0D1B2A",
    "gold":   "C9A84C",
    "header": "154360",
    "light":  "EBF5FB",
    "white":  "FFFFFF",
    "cluster_palette": ["1A5276", "117A65", "B7950B", "922B21", "6C3483", "515A5A"],
}


# ── Data loading ───────────────────────────────────────────────────────────────

def load_leagues(leagues: list[str] | None, min_minutes: int) -> pd.DataFrame:
    if leagues is None:
        paths = sorted(p for p in WYSCOUT_DIR.glob("*.xlsx") if p.stem not in SKIP_FILES)
    else:
        paths = [WYSCOUT_DIR / f"{lg}.xlsx" for lg in leagues]

    frames: list[pd.DataFrame] = []
    for path in paths:
        if not path.exists():
            print(f"  [warn] {path} not found — skipping")
            continue
        lg = path.stem
        try:
            df = pd.read_excel(path)
        except Exception as e:
            print(f"  [warn] Could not read {path.name}: {e}")
            continue
        df = df.copy()
        df["_League"] = lg
        frames.append(df)

    if not frames:
        raise RuntimeError(f"No Wyscout files found in {WYSCOUT_DIR}")

    raw = pd.concat(frames, ignore_index=True)
    print(f"  Loaded {len(frames)} league files, {len(raw)} player rows")

    mins_col = next((c for c in ["Minutes played", "MinutesPlayed", "Minutes"] if c in raw.columns), None)
    raw["_minutes"] = pd.to_numeric(raw[mins_col], errors="coerce").fillna(0) if mins_col else 0

    raw = raw[raw["_minutes"] >= min_minutes].copy()
    print(f"  → {len(raw)} players after {min_minutes}+ minute filter")
    return raw.reset_index(drop=True)


def map_position(pos_str: str) -> str:
    if not isinstance(pos_str, str):
        return "Other"
    first = pos_str.split(",")[0].strip()
    return POS_MAP.get(first, "Other")


def add_position_group(df: pd.DataFrame) -> pd.DataFrame:
    pos_col = next((c for c in ["Position", "Pos"] if c in df.columns), None)
    if pos_col:
        df["_pos_group"] = df[pos_col].apply(map_position)
        df["_full_position"] = df[pos_col].fillna("Unknown")
    else:
        df["_pos_group"] = "Other"
        df["_full_position"] = "Unknown"
    return df


# ── Naming helpers ──────────────────────────────────────────────────────────────

def clean_metric_name(metric: str) -> str:
    name = metric.replace(" per 90", "").replace(".1", "")
    name = re.sub(r",?\s*%$", " %", name)
    return name.strip()


def cluster_label(centroid_z: np.ndarray, metrics: list[str], top_n: int = 2) -> str:
    order = np.argsort(centroid_z)[::-1]
    picks = [i for i in order if centroid_z[i] > 0][:top_n]
    if not picks:
        picks = list(order[:top_n])
    parts = [f"High {clean_metric_name(metrics[i])}" for i in picks]
    return " + ".join(parts)


# ── Clustering ───────────────────────────────────────────────────────────────────

def fit_style_clusters(df_pos: pd.DataFrame, metrics: list[str]) -> dict | None:
    available = [m for m in metrics if m in df_pos.columns]
    if not available:
        return None

    X_raw = df_pos[available].apply(pd.to_numeric, errors="coerce").fillna(0.0).to_numpy()
    n = X_raw.shape[0]
    if n < MIN_POOL_SIZE:
        return None

    scaler = StandardScaler()
    X = scaler.fit_transform(X_raw)

    max_k = min(K_MAX, n // MIN_SAMPLES_PER_K)
    if max_k < K_MIN:
        max_k = K_MIN if n >= K_MIN * MIN_SAMPLES_PER_K // 2 else 2

    best = None
    for k in range(K_MIN, max_k + 1):
        if n < k * MIN_SAMPLES_PER_K:
            continue
        km = KMeans(n_clusters=k, n_init=10, random_state=42)
        labels = km.fit_predict(X)
        if len(set(labels)) < 2:
            continue
        score = silhouette_score(X, labels)
        if best is None or score > best["silhouette"]:
            best = {"k": k, "model": km, "labels": labels, "silhouette": score}

    if best is None:
        km = KMeans(n_clusters=2, n_init=10, random_state=42)
        labels = km.fit_predict(X)
        best = {"k": 2, "model": km, "labels": labels,
                "silhouette": silhouette_score(X, labels) if len(set(labels)) > 1 else 0.0}

    dist = best["model"].transform(X)
    assigned_dist = dist[np.arange(n), best["labels"]]

    return {
        "metrics": available,
        "X_raw": X_raw,
        "labels": best["labels"],
        "centers_z": best["model"].cluster_centers_,
        "k": best["k"],
        "silhouette": best["silhouette"],
        "assigned_dist": assigned_dist,
    }


def style_fit_percentile(fit: dict) -> np.ndarray:
    labels = fit["labels"]
    dist = fit["assigned_dist"]
    out = np.zeros(len(labels))
    for c in np.unique(labels):
        mask = labels == c
        pct = pd.Series(dist[mask]).rank(pct=True) * 100
        out[mask] = 100 - pct.to_numpy()
    return out


def prototype_players(df_pos: pd.DataFrame, fit: dict, cluster_id: int, n: int = 3) -> list[str]:
    labels = fit["labels"]
    dist = fit["assigned_dist"]
    mask = labels == cluster_id
    idx = np.where(mask)[0]
    if len(idx) == 0:
        return []
    order = idx[np.argsort(dist[idx])][:n]
    names = []
    for i in order:
        row = df_pos.iloc[i]
        names.append(f"{row.get('Player', '?')} ({row.get('Team', '?')})")
    return names


# ── Excel helpers ────────────────────────────────────────────────────────────────

def _fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)


def _border() -> Border:
    thin = Side(style="thin", color="CCCCCC")
    return Border(left=thin, right=thin, top=thin, bottom=thin)


def _autofit(ws) -> None:
    for col_cells in ws.columns:
        try:
            max_len = max(
                len(str(col_cells[0].value or "")),
                *(len(str(c.value or "")) for c in col_cells[1:10]),
            )
            ws.column_dimensions[col_cells[0].column_letter].width = min(max_len + 2, 45)
        except Exception:
            pass


def write_data_sheet(ws, title: str, subtitle: str, df: pd.DataFrame, cluster_colors: dict | None = None) -> None:
    ws.append([title])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(len(df.columns), 10))
    ws["A1"].font = Font(bold=True, color=C["white"], size=13)
    ws["A1"].fill = _fill(C["navy"])
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 22

    ws.append([subtitle])
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max(len(df.columns), 10))
    ws["A2"].font = Font(italic=True, color=C["gold"], size=9)
    ws["A2"].fill = _fill(C["navy"])
    ws["A2"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[2].height = 16

    if df.empty:
        return

    ws.append(list(df.columns))
    hdr_row = ws.max_row
    for cell in ws[hdr_row]:
        cell.font = Font(bold=True, color=C["white"], size=9)
        cell.fill = _fill(C["header"])
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _border()
    ws.row_dimensions[hdr_row].height = 28

    cluster_idx = list(df.columns).index("Style Cluster") + 1 if "Style Cluster" in df.columns else None

    for i, row_vals in enumerate(df.itertuples(index=False), start=1):
        ws.append(list(row_vals))
        data_row = ws.max_row
        bg = C["light"] if i % 2 == 0 else C["white"]
        for cell in ws[data_row]:
            cell.font = Font(size=9)
            cell.fill = _fill(bg)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = _border()

        if cluster_idx and cluster_colors:
            cc = ws.cell(data_row, cluster_idx)
            hex_c = cluster_colors.get(cc.value)
            if hex_c:
                cc.fill = _fill(hex_c)
                cc.font = Font(bold=True, color=C["white"], size=9)

    ws.freeze_panes = f"A{hdr_row + 1}"
    _autofit(ws)


def build_readme(ws, leagues: list[str], total: int, min_minutes: int, pos_summaries: dict) -> None:
    ws.title = "README"
    ws.sheet_view.showGridLines = False

    ws.append(["FC HRADEC KRÁLOVÉ — STYLE CLUSTER MODEL"])
    ws.merge_cells("A1:D1")
    ws["A1"].font = Font(bold=True, color=C["white"], size=15)
    ws["A1"].fill = _fill(C["navy"])
    ws.row_dimensions[1].height = 28

    league_label = f"{len(leagues)} leagues" if len(leagues) > 5 else " + ".join(leagues)
    ws.append([f"Waltzing Analytics  ·  Unsupervised playing-style clustering  ·  "
               f"{league_label}  ·  {min_minutes}+ min  ·  {total:,} players"])
    ws.merge_cells("A2:D2")
    ws["A2"].font = Font(italic=True, color=C["gold"], size=10)
    ws["A2"].fill = _fill(C["navy"])
    ws.row_dimensions[2].height = 18

    ws.append([None])
    ws.append([None, "WHAT THIS MODEL DOES"])
    ws[f"B{ws.max_row}"].font = Font(bold=True, size=11, color=C["navy"])

    intro = [
        "Every other model in this workbook set rates players on a QUALITY axis — how good",
        "is this player. This model asks a different question: what KIND of player is this,",
        "independent of how good they are. It clusters each position group on playing-STYLE",
        "metrics only (passing tendencies, crossing, carrying, aerial involvement, zone of",
        "operation) — output metrics like Goals, xG, Assists and xA are deliberately excluded,",
        "since including them would just re-split players into good/bad again.",
        "",
        "Clusters are found with KMeans on standardised style metrics; the number of clusters",
        "per position is chosen by silhouette score, not hand-picked. Cluster labels are",
        "generated directly from the centroid — the two metrics with the highest z-score in",
        "that cluster — so the name is a readout of the data, not a house archetype list.",
    ]
    for line in intro:
        ws.append([None, line])

    ws.append([None])
    ws.append([None, "WORKBOOK STRUCTURE"])
    ws[f"B{ws.max_row}"].font = Font(bold=True, size=11, color=C["navy"])
    ws.append([None, "Sheet", "Contents"])
    for col_letter in "BC":
        cell = ws[f"{col_letter}{ws.max_row}"]
        cell.font = Font(bold=True, color=C["white"])
        cell.fill = _fill(C["header"])
        cell.alignment = Alignment(horizontal="left")

    desc_map = {
        "README":                 "This guide",
        "Style Clusters Overview": "One row per position/cluster — size, label, defining traits, prototype players",
        "GK / CB / FB / DM / CM / W / FW": "Every player in that position group with their assigned style cluster",
    }
    for sheet_name, desc in desc_map.items():
        ws.append([None, sheet_name, desc])
        ws[f"B{ws.max_row}"].font = Font(bold=True, color=C["navy"])

    ws.append([None])
    ws.append([None, "KEY TERMS"])
    ws[f"B{ws.max_row}"].font = Font(bold=True, size=11, color=C["navy"])
    terms = [
        ("Style Cluster",  "Numeric ID of the group this player was assigned to within their position"),
        ("Style Label",    "Auto-generated from the two metrics with the highest z-score in the cluster centroid"),
        ("Style Fit",      "Percentile distance-to-centroid within the cluster (100 = most prototypical example)"),
        ("Silhouette",     "Cluster-quality score for the chosen k (higher = more separated, well-formed clusters)"),
        ("Prototype Players", "The 3 players closest to the cluster centroid — recognisable anchors for the label"),
    ]
    for term, desc in terms:
        ws.append([None, term, desc])
        ws[f"B{ws.max_row}"].font = Font(bold=True)

    ws.append([None])
    ws.append([None, "CLUSTERS FOUND PER POSITION"])
    ws[f"B{ws.max_row}"].font = Font(bold=True, size=11, color=C["navy"])
    ws.append([None, "Position", "k", "Silhouette", "N Players"])
    for cell in ws[ws.max_row]:
        if cell.value:
            cell.font = Font(bold=True, color=C["white"])
            cell.fill = _fill(C["header"])
    for pos, info in pos_summaries.items():
        ws.append([None, pos, info["k"], round(info["silhouette"], 3), info["n"]])
        ws[f"B{ws.max_row}"].font = Font(bold=True)

    ws.append([None])
    ws.append([None, "LIMITATIONS"])
    ws[f"B{ws.max_row}"].font = Font(bold=True, size=11, color=C["navy"])
    caveats = [
        "Clusters are fitted independently per position group and per run — the same player",
        "can land in a differently-numbered cluster if the league selection or minute filter",
        "changes, since the whole standardised feature space shifts. Compare Style Label text,",
        "not raw Cluster IDs, across different runs. Auto-generated labels describe the two",
        "most extreme traits only — read the full metric columns for the complete picture.",
    ]
    for line in caveats:
        ws.append([None, line])

    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 60


def build_overview_sheet(ws, overview_rows: list[dict]) -> None:
    ws.title = "Style Clusters Overview"
    ws.sheet_view.showGridLines = False

    ws.append(["STYLE CLUSTERS OVERVIEW — BY POSITION"])
    ws.merge_cells("A1:H1")
    ws["A1"].font = Font(bold=True, color=C["white"], size=13)
    ws["A1"].fill = _fill(C["navy"])
    ws.row_dimensions[1].height = 22

    ws.append(["One row per position/cluster  ·  Labels auto-generated from the two highest z-score traits in the centroid"])
    ws.merge_cells("A2:H2")
    ws["A2"].font = Font(italic=True, color=C["gold"], size=9)
    ws["A2"].fill = _fill(C["navy"])

    ws.append([None])
    cols = ["Position", "Cluster", "N Players", "Style Label", "Top Trait 1", "Top Trait 2", "Top Trait 3", "Prototype Players"]
    ws.append(cols)
    hdr = ws.max_row
    for cell in ws[hdr]:
        cell.font = Font(bold=True, color=C["white"], size=9)
        cell.fill = _fill(C["header"])
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[hdr].height = 24

    for i, row in enumerate(overview_rows, start=1):
        ws.append([row["pos"], row["cluster"], row["n"], row["label"],
                   row["trait1"], row["trait2"], row["trait3"], row["prototypes"]])
        dr = ws.max_row
        bg = C["light"] if i % 2 == 0 else C["white"]
        for cell in ws[dr]:
            cell.font = Font(size=9)
            cell.fill = _fill(bg)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.cell(dr, 1).font = Font(bold=True, size=9)
        ws.cell(dr, 4).font = Font(bold=True, size=9)

    ws.freeze_panes = f"A{hdr + 1}"
    _autofit(ws)
    ws.column_dimensions["H"].width = 45
    ws.column_dimensions["D"].width = 40


# ── Master table per position ───────────────────────────────────────────────────

def build_position_table(df_pos: pd.DataFrame, fit: dict) -> pd.DataFrame:
    metrics = fit["metrics"]
    fit_pct = style_fit_percentile(fit)

    labels_txt = [cluster_label(fit["centers_z"][c], metrics) for c in range(fit["k"])]

    rows = []
    for i, (_, r) in enumerate(df_pos.iterrows()):
        c = int(fit["labels"][i])
        row = {
            "Player": r.get("Player", ""),
            "Team": r.get("Team", ""),
            "League": r.get("_League", ""),
            "Pos": r.get("_pos_group", ""),
            "Full Position": r.get("_full_position", ""),
            "Age": int(r.get("Age", 0)) if pd.notna(r.get("Age")) else "",
            "Minutes": int(r.get("_minutes", 0)),
            "Style Cluster": c,
            "Style Label": labels_txt[c],
            "Style Fit": round(float(fit_pct[i]), 1),
        }
        for m in metrics:
            val = r.get(m, 0)
            try:
                row[clean_metric_name(m)] = round(float(val), 2) if pd.notna(val) else 0
            except Exception:
                row[clean_metric_name(m)] = 0
        rows.append(row)

    out = pd.DataFrame(rows)
    return out.sort_values(["Style Cluster", "Style Fit"], ascending=[True, False]).reset_index(drop=True)


# ── Main ─────────────────────────────────────────────────────────────────────────

def run(leagues: list[str] | None, min_minutes: int, output: Path) -> None:
    print(f"\n{'='*60}")
    print("  Style Cluster Model Builder")
    print(f"  Leagues: {'ALL' if leagues is None else leagues}")
    print(f"  Min minutes: {min_minutes}")
    print(f"{'='*60}\n")

    print("Loading Wyscout files…")
    raw = load_leagues(leagues, min_minutes)

    raw = add_position_group(raw)
    raw = raw[raw["_pos_group"] != "Other"].copy()
    print(f"  → {len(raw)} players with known position")

    print("Fitting style clusters per position…")
    fits: dict[str, dict] = {}
    tables: dict[str, pd.DataFrame] = {}
    overview_rows: list[dict] = []
    pos_summaries: dict[str, dict] = {}

    for pos, metrics in STYLE_METRICS.items():
        df_pos = raw[raw["_pos_group"] == pos].reset_index(drop=True)
        fit = fit_style_clusters(df_pos, metrics)
        if fit is None:
            print(f"    {pos}: insufficient sample — skipped")
            continue

        fits[pos] = fit
        tables[pos] = build_position_table(df_pos, fit)
        pos_summaries[pos] = {"k": fit["k"], "silhouette": fit["silhouette"], "n": len(df_pos)}
        print(f"    {pos}: n={len(df_pos):>5}  k={fit['k']}  silhouette={fit['silhouette']:.3f}")

        for c in range(fit["k"]):
            mask = fit["labels"] == c
            centroid = fit["centers_z"][c]
            order = np.argsort(centroid)[::-1]
            top3 = [(fit["metrics"][j], centroid[j]) for j in order[:3]]
            trait_txt = [f"{clean_metric_name(m)} (z={z:+.2f})" for m, z in top3]
            protos = prototype_players(df_pos, fit, c)
            overview_rows.append({
                "pos": pos, "cluster": c, "n": int(mask.sum()),
                "label": cluster_label(centroid, fit["metrics"]),
                "trait1": trait_txt[0] if len(trait_txt) > 0 else "",
                "trait2": trait_txt[1] if len(trait_txt) > 1 else "",
                "trait3": trait_txt[2] if len(trait_txt) > 2 else "",
                "prototypes": "; ".join(protos),
            })

    total = sum(len(t) for t in tables.values())
    print(f"  → {total} total players clustered across {len(tables)} position groups")

    print(f"\nWriting workbook → {output}")
    output.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    wb.remove(wb.active)

    print("  Writing README…")
    ws_readme = wb.create_sheet("README")
    league_list = leagues if leagues else sorted(raw["_League"].unique().tolist())
    build_readme(ws_readme, league_list, total, min_minutes, pos_summaries)

    print("  Writing Style Clusters Overview…")
    ws_overview = wb.create_sheet("Style Clusters Overview")
    build_overview_sheet(ws_overview, overview_rows)

    palette = C["cluster_palette"]
    for pos, table in tables.items():
        print(f"  Writing {pos}…")
        n_clusters = fits[pos]["k"]
        cluster_colors = {c: palette[c % len(palette)] for c in range(n_clusters)}
        ws_pos = wb.create_sheet(pos)
        write_data_sheet(
            ws_pos,
            f"{pos} STYLE CLUSTERS — {len(table)} Players, k={n_clusters}",
            f"Silhouette {fits[pos]['silhouette']:.3f}  ·  Sorted by cluster, then Style Fit within cluster",
            table,
            cluster_colors,
        )

    wb.save(output)
    print(f"\n✓ Saved {output}\n")


def main():
    parser = argparse.ArgumentParser(description="Build the Style Cluster Model workbook")
    parser.add_argument("--leagues", nargs="*", default=DEFAULT_LEAGUES,
                        help="League names (Wyscout Files stems). Omit for all leagues.")
    parser.add_argument("--min-minutes", type=int, default=DEFAULT_MIN_MINUTES)
    parser.add_argument("--output", type=Path, default=OUT_DIR / "Style_Cluster_Model.xlsx")
    args = parser.parse_args()

    run(args.leagues, args.min_minutes, args.output)


if __name__ == "__main__":
    main()
